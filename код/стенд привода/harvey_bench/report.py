"""Отчёт стенда в Excel: лист «Результаты» с нормами и итогом, лист «Нагрев» с записью по времени."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

F = "Arial"
THIN = Side(style="thin", color="C9D2D7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR = PatternFill("solid", fgColor="1F2A31")
OK = PatternFill("solid", fgColor="E2EFDA")
BAD = PatternFill("solid", fgColor="F8D7DA")


def _header(ws, row: int, cols: list[str]) -> None:
    for c, h in enumerate(cols, 1):
        cell = ws.cell(row, c, h)
        cell.font = Font(name=F, bold=True, color="FFFFFF")
        cell.fill = HDR
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def write(path: Path, results: list[dict], heat_log: list[list]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Результаты"
    ws["A1"] = "Харви — стенд привода STS3215-C018"
    ws["A1"].font = Font(name=F, bold=True, size=14)
    _header(ws, 3, ["№", "Замер", "Измерено", "Норма", "Итог", "Примечание"])
    for i, r in enumerate(results, 4):
        vals = [r["n"], r["test"], r["measured"], r["norm"], r["verdict"], r["note"]]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(i, c, v)
            cell.font = Font(name=F, size=10, bold=(c == 5))
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if r["verdict"] == "ОК":
            ws.cell(i, 5).fill = OK
        elif r["verdict"] == "НЕ ОК":
            ws.cell(i, 5).fill = BAD
    for col, w in zip("ABCDEF", (6, 34, 34, 30, 9, 50)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"

    lg = wb.create_sheet("Нагрев")
    _header(lg, 1, ["Время, с", "Температура, °C", "Ток, А", "Напряжение, В", "Позиция"])
    for i, row in enumerate(heat_log, 2):
        for c, v in enumerate(row, 1):
            lg.cell(i, c, v).font = Font(name=F, size=10)
    for col in "ABCDE":
        lg.column_dimensions[col].width = 16
    lg.freeze_panes = "A2"
    wb.save(path)
